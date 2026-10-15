"""Export chats, notes, reviews, tables and references to Markdown, Word, PDF, LaTeX, HTML, text, CSV, Excel...

Citations in app content are numbered per answer ([1] refers to that answer's source list). On export they are
mapped to papers and re-rendered in the chosen mode:
  numbered     [1], [2, 5]        with a numbered reference list
  author-year  (Vaswani et al., 2017, p. 5)  with an alphabetical reference list
  none         citations removed, no reference list
"""
import csv
import html as htmlmod
import io
import json
import re
import zipfile
from dataclasses import dataclass, field

FORMATS = {
    "Markdown (.md)": "md",
    "Word (.docx)": "docx",
    "PDF (.pdf)": "pdf",
    "LaTeX (.tex, Overleaf-ready)": "tex",
    "HTML (.html)": "html",
    "Plain text (.txt)": "txt",
}
CITATION_MODES = {
    "Numbered [1]": "numbered",
    "Author–year (Smith et al., 2020)": "author-year",
    "No citations (clean text)": "none",
}
STYLES = ["APA 7", "IEEE", "Harvard", "MLA 9", "Chicago (author-date)", "Vancouver"]
NUMERIC_STYLES = {"IEEE", "Vancouver"}
MIME = {
    "md": "text/markdown", "txt": "text/plain", "html": "text/html", "tex": "application/x-tex",
    "pdf": "application/pdf", "bib": "application/x-bibtex", "ris": "application/x-research-info-systems",
    "csv": "text/csv", "json": "application/json", "zip": "application/zip",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

_TOKEN = re.compile(r"⟦c:([^⟧]*)⟧")
_NUM_CITE = re.compile(r"(?:\s*\[\d{1,3}(?:\s*[,–-]\s*\d{1,3})*\])+")


# ---- references ---------------------------------------------------------------------

def _names(doc) -> list[str]:
    return [a.strip() for a in (doc.authors or "").split(";") if a.strip()]


def _split_name(name: str) -> tuple[str, str]:
    """('Vaswani', 'A.') from 'Ashish Vaswani' or 'Vaswani, Ashish'."""
    if "," in name:
        last, first = [x.strip() for x in name.split(",", 1)]
    else:
        parts = name.split()
        last, first = (parts[-1], " ".join(parts[:-1])) if len(parts) > 1 else (name, "")
    initials = " ".join(f"{p[0]}." for p in re.split(r"[\s\-]+", first) if p and p[0].isalpha())
    return last, initials


def short_cite(doc) -> str:
    """'Vaswani et al., 2017' style author-year label."""
    names = _names(doc)
    year = str(doc.year or "n.d.")
    if not names:
        return f"{doc.title[:40]}, {year}"
    lasts = [_split_name(n)[0] for n in names]
    who = lasts[0] if len(lasts) == 1 else f"{lasts[0]} & {lasts[1]}" if len(lasts) == 2 else f"{lasts[0]} et al."
    return f"{who}, {year}"


def format_reference(doc, style: str, n: int | None = None) -> str:
    """One reference-list entry in Markdown (italics with *...*)."""
    names = _names(doc)
    parts = [_split_name(x) for x in names]
    year = doc.year or "n.d."
    title = doc.title.rstrip(".")
    venue = doc.venue or ("arXiv" if doc.arxiv_id else "")
    link = f"https://doi.org/{doc.doi}" if doc.doi else (doc.url or "")
    if style == "IEEE":
        au = [f"{i} {l}".strip() for l, i in parts]
        au_s = au[0] if len(au) == 1 else f"{au[0]} and {au[1]}" if len(au) == 2 else \
            (", ".join(au[:-1]) + ", and " + au[-1]) if len(au) <= 6 else (f"{au[0]} et al." if au else "")
        s = f"{au_s}, “{title},” *{venue}*" if venue else f"{au_s}, “{title}”"
        s += f", {year}." + (f" doi: {doc.doi}." if doc.doi else (f" [Online]. Available: {link}" if link else ""))
        return f"[{n}] {s}" if n else s
    if style == "Vancouver":
        au = [f"{l} {i.replace('.', '').replace(' ', '')}".strip() for l, i in parts]
        au_s = ", ".join(au[:6]) + (", et al" if len(au) > 6 else "")
        s = f"{au_s}. {title}. {venue + '. ' if venue else ''}{year}." + (f" doi:{doc.doi}" if doc.doi else "")
        return f"{n}. {s}" if n else s
    if style == "MLA 9":
        if not parts:
            au_s = ""
        elif len(parts) == 1:
            au_s = f"{parts[0][0]}, {names[0].split()[0] if ' ' in names[0] else ''}".rstrip(", ") + ". "
        elif len(parts) == 2:
            au_s = f"{parts[0][0]}, {names[0].split()[0]}, and {names[1]}. "
        else:
            au_s = f"{parts[0][0]}, {names[0].split()[0]}, et al. "
        return f"{au_s}“{title}.” *{venue}*, {year}" + (f", {link}." if link else ".")
    # Author-date styles: APA, Harvard, Chicago
    if style.startswith("APA"):
        au = [f"{l}, {i}".strip(", ") for l, i in parts]
        au_s = au[0] if len(au) == 1 else (", ".join(au[:-1]) + ", & " + au[-1]) if 1 < len(au) <= 20 else \
            ", ".join(au[:19]) + ", … " + au[-1]
        return f"{au_s + ' ' if au_s else ''}({year}). {title}. " + (f"*{venue}*. " if venue else "") + (link or "")
    if style.startswith("Chicago"):
        au = [f"{l}, {names[0].split()[0]}" if k == 0 and " " in names[0] else n_ for k, (n_, (l, i)) in
              enumerate(zip(names, parts))]
        au_s = au[0] if len(au) == 1 else (", ".join(au[:-1]) + ", and " + au[-1]) if au else ""
        return f"{au_s + '. ' if au_s else ''}{year}. “{title}.” *{venue}*." + (f" {link}." if link else "")
    # Harvard
    au = [f"{l}, {i}".strip(", ") for l, i in parts]
    au_s = au[0] if len(au) == 1 else (", ".join(au[:-1]) + " and " + au[-1]) if au else ""
    return f"{au_s + ' ' if au_s else ''}({year}) ‘{title}’, *{venue}*." + (f" Available at: {link}." if link else "")


def bib_keys(docs) -> dict[int, str]:
    """Unique, stable BibTeX keys like vaswani2017attention (a/b suffixes on clashes)."""
    keys, used = {}, {}
    for d in docs:
        names = _names(d)
        last = re.sub(r"\W", "", _split_name(names[0])[0]).lower() if names else "anon"
        word = next((w.lower() for w in re.findall(r"[A-Za-z]+", d.title) if len(w) > 3), "paper")
        base = f"{last}{d.year or ''}{word}"
        n = used.get(base, 0)
        used[base] = n + 1
        keys[d.id] = base if n == 0 else f"{base}{chr(ord('a') + n)}"
    return keys


def _bib_escape(s: str) -> str:
    return (s or "").replace("&", r"\&").replace("%", r"\%").replace("#", r"\#")


def bibtex(docs) -> str:
    keys = bib_keys(docs)
    out = []
    for d in docs:
        is_preprint = (d.venue or "").lower() in ("", "arxiv")
        kind = "misc" if is_preprint else "article"
        fields = {
            "title": "{" + _bib_escape(d.title) + "}",
            "author": " and ".join(_names(d)),
            "year": str(d.year or ""),
            ("journal" if kind == "article" else "howpublished"):
                _bib_escape(d.venue) if kind == "article" else (f"arXiv:{d.arxiv_id}" if d.arxiv_id else ""),
            "doi": d.doi, "eprint": d.arxiv_id, "archivePrefix": "arXiv" if d.arxiv_id else "", "url": d.url,
        }
        body = ",\n".join(f"  {k} = {{{v}}}" for k, v in fields.items() if v)
        out.append(f"@{kind}{{{keys[d.id]},\n{body}\n}}")
    return "\n\n".join(out) + "\n"


def ris(docs) -> str:
    out = []
    for d in docs:
        lines = ["TY  - JOUR" if d.venue and d.venue.lower() != "arxiv" else "TY  - GEN", f"TI  - {d.title}"]
        lines += [f"AU  - {a}" for a in _names(d)]
        if d.year:
            lines.append(f"PY  - {d.year}")
        if d.venue:
            lines.append(f"T2  - {d.venue}")
        if d.doi:
            lines.append(f"DO  - {d.doi}")
        if d.url:
            lines.append(f"UR  - {d.url}")
        if d.abstract:
            lines.append(f"AB  - {d.abstract}")
        if d.tags:
            lines += [f"KW  - {t.strip()}" for t in d.tags.split(",") if t.strip()]
        lines.append("ER  - ")
        out.append("\n".join(lines))
    return "\n\n".join(out) + "\n"


def references_csv(docs) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Title", "Authors", "Year", "Venue", "DOI", "arXiv", "URL", "Status", "Tags", "Abstract"])
    for d in docs:
        w.writerow([d.title, d.authors, d.year or "", d.venue, d.doi, d.arxiv_id, d.url, d.status, d.tags, d.abstract])
    return buf.getvalue()


# ---- citation handling ----------------------------------------------------------------

def tokenize_citations(text: str, sources: list[dict]) -> str:
    """Replace per-answer [n] markers with paper tokens ⟦c:doc@page;doc@page⟧ using that answer's sources."""
    def repl(m):
        nums = []
        for grp in re.findall(r"\[([^\]]+)\]", m.group(0)):
            for part in re.split(r"\s*,\s*", grp):
                if re.fullmatch(r"\d+\s*[–-]\s*\d+", part):
                    a, b = [int(x) for x in re.split(r"\s*[–-]\s*", part)]
                    nums += list(range(a, b + 1)) if 0 < b - a < 20 else [a, b]
                elif part.strip().isdigit():
                    nums.append(int(part))
        refs = []
        for n in nums:
            if 0 < n <= len(sources) and sources[n - 1].get("doc_id"):
                s = sources[n - 1]
                item = f"{s['doc_id']}@{s.get('page') or ''}"
                if item not in refs:
                    refs.append(item)
        return f" ⟦c:{';'.join(refs)}⟧" if refs else m.group(0)
    return _NUM_CITE.sub(repl, text)


@dataclass
class Citations:
    docs: dict  # doc_id -> Document
    mode: str = "numbered"
    style: str = "APA 7"
    order: list = field(default_factory=list)  # doc ids in order of first citation

    def __post_init__(self):
        self.keys = bib_keys(list(self.docs.values()))

    def _items(self, payload: str) -> list[tuple[int, str]]:
        out = []
        for it in payload.split(";"):
            if "@" in it:
                d, p = it.split("@", 1)
                if d.isdigit() and int(d) in self.docs:
                    out.append((int(d), p))
                    if int(d) not in self.order:
                        self.order.append(int(d))
        return out

    def inline(self, payload: str, fmt: str) -> str:
        items = self._items(payload)
        if not items or self.mode == "none":
            return ""
        if fmt == "tex":
            keys = list(dict.fromkeys(self.keys[d] for d, _ in items))
            page = items[0][1] if len(keys) == 1 and items[0][1] else ""
            cmd = "\\citep" if self.mode == "author-year" else "\\cite"
            return f"{cmd}[p.~{page}]{{{keys[0]}}}" if page else f"{cmd}{{{','.join(keys)}}}"
        if self.mode == "author-year":
            pages: dict[int, list[str]] = {}
            for d, p in items:
                pages.setdefault(d, [])
                if p and p not in pages[d]:
                    pages[d].append(p)
            parts = []
            for d, ps in pages.items():
                ps = sorted(ps, key=lambda x: int(x) if x.isdigit() else 0)
                where = f", p. {ps[0]}" if len(ps) == 1 else (f", pp. {', '.join(ps)}" if ps else "")
                parts.append(short_cite(self.docs[d]) + where)
            return "(" + "; ".join(parts) + ")"
        nums = sorted({self.order.index(d) + 1 for d, _ in items})
        return "[" + ", ".join(str(n) for n in nums) + "]"

    def strip(self, text: str, fmt: str) -> str:
        """Render all tokens in already-converted text."""
        out = _TOKEN.sub(lambda m: self.inline(m.group(1), fmt), text)
        if self.mode == "none":
            out = re.sub(r"[ \t]+([.,;:)])", r"\1", out)
            out = re.sub(r"[ \t]+(\n|\||$)", r"\1", out)
        return out

    def reference_docs(self) -> list:
        docs = [self.docs[d] for d in self.order if d in self.docs]
        if self.mode == "author-year" or self.style not in NUMERIC_STYLES:
            docs.sort(key=lambda d: (short_cite(d).lower(), d.year or 0))
        return docs

    def reference_list(self) -> list[str]:
        """Formatted entries: in citation order with numbers, or alphabetical for author-year."""
        if self.mode == "none":
            return []
        if self.mode == "author-year":
            return [format_reference(d, self.style) for d in self.reference_docs()]
        cited = [self.docs[x] for x in self.order if x in self.docs]
        if self.style in NUMERIC_STYLES:
            return [format_reference(d, self.style, i) for i, d in enumerate(cited, 1)]
        return [f"[{i}] {format_reference(d, self.style)}" for i, d in enumerate(cited, 1)]


# ---- tiny Markdown parser shared by the LaTeX / Word / text renderers --------------------------

def parse_blocks(md: str) -> list[tuple]:
    """Blocks: ('h', level, text) ('p', text) ('ul'|'ol', [(indent, text)]) ('table', rows) ('code', text) ('quote', text) ('hr',)."""
    lines = md.replace("\r\n", "\n").split("\n")
    blocks, i = [], 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("```"):
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith("```"):
                j += 1
            blocks.append(("code", "\n".join(lines[i + 1:j])))
            i = j + 1
            continue
        if m := re.match(r"^(#{1,6})\s+(.*)", line):
            blocks.append(("h", len(m.group(1)), m.group(2).strip().strip("#").strip()))
            i += 1
            continue
        if re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", line):
            blocks.append(("hr",))
            i += 1
            continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\s*\|?[\s:\-|]+\|[\s:\-|]*$", lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append(("table", rows))
            continue
        if re.match(r"^\s*([-*+]|\d+[.)])\s+", line):
            kind = "ol" if re.match(r"^\s*\d+[.)]\s+", line) else "ul"
            items = []
            while i < len(lines) and (m := re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)", lines[i])):
                items.append((len(m.group(1).replace("\t", "    ")) // 2, m.group(3)))
                i += 1
                while i < len(lines) and lines[i].startswith("   ") and lines[i].strip() and \
                        not re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[i]):
                    items[-1] = (items[-1][0], items[-1][1] + " " + lines[i].strip())
                    i += 1
            blocks.append((kind, items))
            continue
        if line.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].startswith(">"):
                buf.append(lines[i].lstrip("> "))
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue
        buf = []
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|```|\s*([-*+]|\d+[.)])\s+|\|)", lines[i]):
            buf.append(lines[i].strip())
            i += 1
        if buf:
            blocks.append(("p", " ".join(buf)))
        else:
            i += 1
    return blocks


_INLINE = re.compile(r"(⟦c:[^⟧]*⟧|\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|:[a-z]+-badge\[[^\]]*\])")


def inline_parts(text: str) -> list[tuple[str, str]]:
    """Split inline Markdown into (kind, text): plain, bold, italic, code, link(text|url), cite(payload)."""
    out, pos = [], 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            out.append(("plain", text[pos:m.start()]))
        t = m.group(0)
        if t.startswith("⟦"):
            out.append(("cite", t[3:-1]))
        elif t.startswith(("**", "__")):
            out.append(("bold", t[2:-2]))
        elif t.startswith("`"):
            out.append(("code", t[1:-1]))
        elif t.startswith("["):
            lm = re.match(r"\[([^\]]+)\]\(([^)]+)\)", t)
            out.append(("link", f"{lm.group(1)}|{lm.group(2)}"))
        elif t.startswith(":"):
            out.append(("plain", re.search(r"\[([^\]]*)\]", t).group(1)))
        else:
            out.append(("italic", t[1:-1]))
        pos = m.end()
    if pos < len(text):
        out.append(("plain", text[pos:]))
    return out


# ---- LaTeX ----------------------------------------------------------------------------------

_TEX_ESC = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
            "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


# Characters pdflatex can't typeset from UTF-8 on its own (common in papers: Greek letters, math symbols).
_TEX_UNICODE = {
    "“": "``", "”": "''", "‘": "`", "’": "'", "–": "--", "—": "---", "…": r"\ldots{}", "°": r"\textdegree{}",
    "≥": r"$\geq$", "≤": r"$\leq$", "±": r"$\pm$", "×": r"$\times$", "→": r"$\rightarrow$", "←": r"$\leftarrow$",
    "≈": r"$\approx$", "∼": r"$\sim$", "≠": r"$\neq$", "∞": r"$\infty$", "√": r"$\surd$", "∈": r"$\in$",
    "·": r"$\cdot$", "−": "-", "µ": r"$\mu$", "²": r"$^2$", "³": r"$^3$", "€": "EUR", "•": r"\textbullet{}",
    "α": r"$\alpha$", "β": r"$\beta$", "γ": r"$\gamma$", "δ": r"$\delta$", "ε": r"$\epsilon$", "ζ": r"$\zeta$",
    "η": r"$\eta$", "θ": r"$\theta$", "κ": r"$\kappa$", "λ": r"$\lambda$", "μ": r"$\mu$", "ν": r"$\nu$",
    "ξ": r"$\xi$", "π": r"$\pi$", "ρ": r"$\rho$", "σ": r"$\sigma$", "τ": r"$\tau$", "φ": r"$\phi$",
    "χ": r"$\chi$", "ψ": r"$\psi$", "ω": r"$\omega$", "Γ": r"$\Gamma$", "Δ": r"$\Delta$", "Θ": r"$\Theta$",
    "Λ": r"$\Lambda$", "Π": r"$\Pi$", "Σ": r"$\Sigma$", "Φ": r"$\Phi$", "Ψ": r"$\Psi$", "Ω": r"$\Omega$",
}


def tex_escape(s: str) -> str:
    s = "".join(_TEX_ESC.get(ch, ch) for ch in s)
    s = "".join(_TEX_UNICODE.get(ch, ch) for ch in s)
    # Drop emoji and other symbols outside the Latin/Greek ranges that pdflatex can't render.
    return "".join(ch for ch in s if ord(ch) < 0x2000 or ch in "€")


def tex_inline(text: str, cites: Citations) -> str:
    out = []
    for kind, t in inline_parts(text):
        if kind == "cite":
            c = cites.inline(t, "tex")
            out.append(("~" + c) if c else "")
        elif kind == "bold":
            out.append(r"\textbf{" + tex_escape(t) + "}")
        elif kind == "italic":
            out.append(r"\emph{" + tex_escape(t) + "}")
        elif kind == "code":
            out.append(r"\texttt{" + tex_escape(t) + "}")
        elif kind == "link":
            label, url = t.split("|", 1)
            out.append(r"\href{" + url.replace("%", r"\%").replace("#", r"\#") + "}{" + tex_escape(label) + "}")
        else:
            out.append(tex_escape(t))
    return re.sub(r"\s+~", "~", "".join(out))


def tex_body(md: str, cites: Citations, base_level: int = 1) -> str:
    cmds = ["section", "subsection", "subsubsection", "paragraph", "paragraph", "paragraph"]
    out = []
    for b in parse_blocks(md):
        if b[0] == "h":
            out.append(f"\\{cmds[min(5, b[1] - 1 + base_level - 1)]}*{{{tex_inline(b[2], cites)}}}")
        elif b[0] == "p":
            out.append(tex_inline(b[1], cites))
        elif b[0] in ("ul", "ol"):
            env = "itemize" if b[0] == "ul" else "enumerate"
            depth, lines = 0, []
            for indent, t in b[1]:
                indent = min(indent, 3)
                while depth < indent:
                    lines.append(f"\\begin{{{env}}}")
                    depth += 1
                while depth > indent:
                    lines.append(f"\\end{{{env}}}")
                    depth -= 1
                lines.append(r"  \item " + tex_inline(t, cites))
            lines += [f"\\end{{{env}}}"] * depth
            out.append(f"\\begin{{{env}}}\n" + "\n".join(lines) + f"\n\\end{{{env}}}")
        elif b[0] == "table" and b[1]:
            ncol = max(len(r) for r in b[1])
            width = max(0.08, 0.95 / ncol)
            spec = "".join(f"p{{{width:.2f}\\linewidth}}" for _ in range(ncol))
            rows = [" & ".join(tex_inline(c, cites) for c in (r + [""] * (ncol - len(r)))) + r" \\" for r in b[1]]
            out.append("{\\small\n\\begin{longtable}{" + spec + "}\n\\toprule\n" + rows[0] + "\n\\midrule\n\\endhead\n"
                       + "\n".join(rows[1:]) + "\n\\bottomrule\n\\end{longtable}\n}")
        elif b[0] == "code":
            out.append("\\begin{verbatim}\n" + b[1] + "\n\\end{verbatim}")
        elif b[0] == "quote":
            out.append("\\begin{quote}\n" + tex_inline(b[1], cites) + "\n\\end{quote}")
        elif b[0] == "hr":
            out.append(r"\noindent\rule{\linewidth}{0.4pt}")
    return "\n\n".join(out)


def tex_document(title: str, body: str, cites: Citations, landscape: bool = False, author: str = "") -> str:
    natbib = "\\usepackage[numbers,sort&compress]{natbib}" if cites.mode != "author-year" else \
        "\\usepackage[authoryear,round]{natbib}"
    bibstyle = "plainnat" if cites.mode == "author-year" else ("ieeetr" if cites.style == "IEEE" else "unsrtnat")
    used = [cites.docs[d] for d in cites.order if d in cites.docs]
    bib = ""
    if cites.mode != "none" and used:
        bib = ("\n\\bibliographystyle{" + bibstyle + "}\n\\bibliography{references}\n")
    filecontents = ("\\begin{filecontents*}[overwrite]{references.bib}\n" + bibtex(used) + "\\end{filecontents*}\n"
                    if bib else "")
    return (
        filecontents
        + "\\documentclass[11pt,a4paper" + (",landscape" if landscape else "") + "]{article}\n"
        "\\usepackage[utf8]{inputenc}\n\\usepackage[T1]{fontenc}\n\\usepackage{textcomp}\n\\usepackage{lmodern}\n"
        "\\usepackage[margin=2.2cm]{geometry}\n\\usepackage{booktabs}\n\\usepackage{longtable}\n"
        "\\usepackage{array}\n\\usepackage{enumitem}\n\\usepackage[hidelinks]{hyperref}\n" + natbib + "\n\n"
        f"\\title{{{tex_escape(title)}}}\n\\author{{{tex_escape(author)}}}\n\\date{{\\today}}\n\n"
        "\\begin{document}\n\\maketitle\n\n" + body + "\n" + bib + "\\end{document}\n"
    )


# ---- Word -------------------------------------------------------------------------------------

def _docx_inline(par, text: str, cites: Citations) -> None:
    for kind, t in inline_parts(text):
        if kind == "cite":
            c = cites.inline(t, "docx")
            if c:
                par.add_run(c)  # the token already carries its leading space
        elif kind == "link":
            label, url = t.split("|", 1)
            r = par.add_run(label)
            r.underline = True
        else:
            r = par.add_run(t)
            r.bold = kind == "bold"
            r.italic = kind == "italic"
            if kind == "code":
                r.font.name = "Courier New"


def docx_add_markdown(doc, md: str, cites: Citations, base_level: int = 1) -> None:
    for b in parse_blocks(md):
        if b[0] == "h":
            h = doc.add_heading(level=min(9, b[1] + base_level - 1))
            _docx_inline(h, b[2], cites)
        elif b[0] == "p":
            _docx_inline(doc.add_paragraph(), b[1], cites)
        elif b[0] in ("ul", "ol"):
            for indent, t in b[1]:
                style = ("List Bullet" if b[0] == "ul" else "List Number") + (f" {min(3, indent + 1)}" if indent else "")
                try:
                    p = doc.add_paragraph(style=style)
                except KeyError:
                    p = doc.add_paragraph(style="List Bullet")
                _docx_inline(p, t, cites)
        elif b[0] == "table" and b[1]:
            ncol = max(len(r) for r in b[1])
            table = doc.add_table(rows=len(b[1]), cols=ncol)
            table.style = "Light Grid Accent 1"
            for r_i, row in enumerate(b[1]):
                for c_i in range(ncol):
                    cell = table.cell(r_i, c_i)
                    cell.text = ""
                    _docx_inline(cell.paragraphs[0], row[c_i] if c_i < len(row) else "", cites)
                    if r_i == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            doc.add_paragraph()
        elif b[0] == "code":
            p = doc.add_paragraph()
            r = p.add_run(b[1])
            r.font.name = "Courier New"
        elif b[0] == "quote":
            p = doc.add_paragraph(style="Intense Quote")
            _docx_inline(p, b[1], cites)


def new_docx(title: str):
    from docx import Document as Docx
    from docx.shared import Pt
    doc = Docx()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, level=0)
    return doc


def docx_bytes(doc) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---- HTML / PDF / text ----------------------------------------------------------------------------

HTML_CSS = """
body{font-family:Georgia,'Times New Roman',serif;max-width:820px;margin:40px auto;padding:0 20px;line-height:1.55;color:#1f2328}
h1,h2,h3,h4{font-family:Helvetica,Arial,sans-serif;line-height:1.25}
table{border-collapse:collapse;width:100%;margin:12px 0;font-size:0.9em}
th,td{border:1px solid #ccc;padding:6px 8px;vertical-align:top;text-align:left}
th{background:#f3f4f6}
.refs li{margin-bottom:6px}
.meta{color:#666;font-size:0.9em}
blockquote{border-left:3px solid #ccc;margin:0;padding-left:12px;color:#555}
"""


def md_render_citations(md: str, cites: Citations) -> str:
    """Markdown with tokens rendered as plain citations (for .md, HTML, PDF and text)."""
    return cites.strip(md, "md")


def html_from_md(md: str) -> str:
    import markdown
    return markdown.markdown(md, extensions=["tables", "sane_lists"])


def html_document(title: str, body_html: str) -> str:
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{htmlmod.escape(title)}</title>"
            f"<style>{HTML_CSS}</style></head><body><h1>{htmlmod.escape(title)}</h1>{body_html}</body></html>")


def pdf_from_html(body_html: str, title: str, landscape: bool = False) -> bytes:
    import pymupdf
    css = ("body{font-family:sans-serif;font-size:10pt;line-height:1.4} h1{font-size:18pt} h2{font-size:14pt} "
           "h3{font-size:12pt} table{border-collapse:collapse} th,td{border:1px solid #999;padding:3px;font-size:8.5pt} "
           "th{background-color:#eeeeee} li{margin-bottom:3px}")
    story = pymupdf.Story(html=f"<h1>{htmlmod.escape(title)}</h1>{body_html}", user_css=css)
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    box = pymupdf.paper_rect("a4-l" if landscape else "a4")
    where = box + (45, 50, -45, -50)
    more = True
    while more:
        dev = writer.begin_page(box)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return buf.getvalue()


def text_from_md(md: str) -> str:
    out = []
    for b in parse_blocks(md):
        flat = lambda t: "".join(x if k != "link" else x.split("|")[0] for k, x in inline_parts(t) if k != "cite")
        if b[0] == "h":
            t = flat(b[2])
            out.append(t.upper() if b[1] <= 2 else t)
        elif b[0] == "p":
            out.append(flat(b[1]))
        elif b[0] in ("ul", "ol"):
            out.append("\n".join(f"{'  ' * ind}{'-' if b[0] == 'ul' else f'{k}.'} {flat(t)}"
                                 for k, (ind, t) in enumerate(b[1], 1)))
        elif b[0] == "table":
            out.append("\n".join(" | ".join(flat(c) for c in r) for r in b[1]))
        elif b[0] in ("code", "quote"):
            out.append(b[1])
    return "\n\n".join(out) + "\n"


# ---- document export (chat, review, notes) ------------------------------------------------------------

@dataclass
class Section:
    heading: str        # may be empty
    body: str           # Markdown with per-section [n] citations
    sources: list = field(default_factory=list)  # that section's numbered sources (dicts with doc_id, page)
    level: int = 2


def export_document(title: str, sections: list[Section], docs: dict, fmt: str, mode: str, style: str,
                    subtitle: str = "") -> bytes:
    """Render sections into one document with a consistent, global reference list."""
    cites = Citations(docs, mode, style)
    bodies = [(s, tokenize_citations(s.body, s.sources) if s.sources else s.body) for s in sections]
    refs_heading = "References"

    if fmt == "tex":
        parts = []
        if subtitle:
            parts.append(r"\noindent " + tex_escape(subtitle) + "\n")
        for s, body in bodies:
            if s.heading:
                cmd = "section" if s.level <= 2 else "subsection"
                parts.append(f"\\{cmd}*{{{tex_inline(s.heading, cites)}}}")
            parts.append(tex_body(body, cites, base_level=2 if s.heading else 1))
        return tex_document(title, "\n\n".join(parts), cites).encode()

    if fmt == "docx":
        d = new_docx(title)
        if subtitle:
            d.add_paragraph(subtitle).italic = True
        for s, body in bodies:
            if s.heading:
                h = d.add_heading(level=1 if s.level <= 2 else 2)
                _docx_inline(h, s.heading, cites)
            docx_add_markdown(d, body, cites, base_level=2 if s.heading else 1)
        refs = cites.reference_list()
        if refs:
            d.add_heading(refs_heading, level=1)
            for r in refs:
                p = d.add_paragraph()
                _docx_inline(p, r, Citations({}, "none"))
        return docx_bytes(d)

    # Markdown-based formats
    md_parts = [f"*{subtitle}*\n"] if subtitle else []
    for s, body in bodies:
        if s.heading:
            md_parts.append(f"{'#' * s.level} {s.heading}\n")
        md_parts.append(cites.strip(body, "md").strip() + "\n")
    refs = cites.reference_list()
    if refs:
        md_parts.append(f"## {refs_heading}\n")
        md_parts.append("\n\n".join(refs))
    md = "\n".join(md_parts)
    if fmt == "md":
        return (f"# {title}\n\n" + md).encode()
    if fmt == "txt":
        return (title.upper() + "\n\n" + text_from_md(md)).encode()
    body_html = html_from_md(md)
    if fmt == "html":
        return html_document(title, body_html).encode()
    if fmt == "pdf":
        return pdf_from_html(body_html, title)
    raise ValueError(fmt)


# ---- table export -------------------------------------------------------------------------------------

def table_rows(table: dict, docs: dict, cells: dict, with_citations: bool, cites: Citations | None = None):
    """Header + rows for an extraction table. Citation mode adds a reference/key column."""
    cols = [c["name"] for c in table["columns"]]
    header = ["Paper", "Authors", "Year", "Venue", "DOI"] + (["Citation"] if with_citations else []) + cols
    rows = []
    keys = bib_keys([docs[d] for d in table["doc_ids"] if d in docs])
    for doc_id in table["doc_ids"]:
        d = docs.get(doc_id)
        if not d:
            continue
        row = [d.title, d.authors.replace(";", ","), d.year or "", d.venue, d.doi]
        if with_citations:
            row.append(short_cite(d) if not cites else keys[doc_id])
        row += [cells.get((doc_id, c), {}).get("value", "") for c in cols]
        rows.append((doc_id, row))
    return header, rows


def evidence_rows(table: dict, docs: dict, cells: dict) -> list[list]:
    out = []
    for doc_id in table["doc_ids"]:
        d = docs.get(doc_id)
        if not d:
            continue
        for c in table["columns"]:
            cell = cells.get((doc_id, c["name"]), {})
            srcs = cell.get("sources") or [{}]
            for s in srcs:
                out.append([d.title, short_cite(d), c["name"], cell.get("value", ""), s.get("page", ""),
                            (s.get("text") or "")[:1500]])
    return out


def export_table(table: dict, docs: dict, cells: dict, fmt: str, with_citations: bool, style: str = "APA 7") -> bytes:
    used = [docs[d] for d in table["doc_ids"] if d in docs]
    header, rows = table_rows(table, docs, cells, with_citations)
    name = table.get("name") or "Extraction table"

    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(r for _, r in rows)
        return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly

    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Extraction"
        ws.append(header)
        for _, r in rows:
            ws.append(r)
        sheets = [ws]
        if with_citations:
            ev = wb.create_sheet("Evidence")
            ev.append(["Paper", "Citation", "Column", "Value", "Page", "Supporting quote"])
            for r in evidence_rows(table, docs, cells):
                ev.append(r)
            ref = wb.create_sheet("References")
            ref.append([f"Reference ({style})"])
            for d in sorted(used, key=lambda d: short_cite(d).lower()):
                ref.append([re.sub(r"\*", "", format_reference(d, style))])
            sheets += [ev, ref]
        for sh in sheets:
            for cell in sh[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="4F46E5")
            for col in sh.columns:
                width = min(60, max(10, *(len(str(c.value or "")) for c in col[:50])))
                sh.column_dimensions[col[0].column_letter].width = width
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
            sh.freeze_panes = "B2"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    if fmt == "json":
        out = []
        for doc_id, r in rows:
            item = dict(zip(header, r))
            if with_citations:
                item["evidence"] = {c["name"]: cells.get((doc_id, c["name"]), {}).get("sources", [])
                                    for c in table["columns"]}
            out.append(item)
        return json.dumps(out, indent=2, ensure_ascii=False).encode()

    # Compact view for documents: drop Authors/Venue/DOI, keep a citation column.
    cols = [c["name"] for c in table["columns"]]
    if fmt == "tex":
        cites = Citations({d.id: d for d in used}, "numbered" if with_citations else "none")
        keys = bib_keys(used)
        ncol = 2 + len(cols)
        width = max(0.07, 0.93 / ncol)
        spec = f"p{{{min(0.2, width * 1.6):.2f}\\linewidth}}p{{0.05\\linewidth}}" + \
               "".join(f"p{{{width:.2f}\\linewidth}}" for _ in cols)
        head = " & ".join([r"\textbf{Paper}", r"\textbf{Year}"] + [r"\textbf{" + tex_escape(c) + "}" for c in cols]) + r" \\"
        lines = []
        for doc_id, r in rows:
            d = docs[doc_id]
            paper = tex_escape(short_cite(d).rsplit(",", 1)[0]) + (f"~\\cite{{{keys[doc_id]}}}" if with_citations else "")
            if with_citations:
                cites.order.append(doc_id)
            vals = [tex_escape(str(cells.get((doc_id, c), {}).get("value", ""))) for c in cols]
            lines.append(" & ".join([paper, str(d.year or "")] + vals) + r" \\")
        body = ("{\\scriptsize\n\\setlength{\\tabcolsep}{3pt}\n\\begin{longtable}{" + spec + "}\n"
                f"\\caption{{{tex_escape(name)}}}\\\\\n\\toprule\n{head}\n\\midrule\n\\endfirsthead\n\\toprule\n{head}\n"
                "\\midrule\n\\endhead\n" + "\n".join(lines) + "\n\\bottomrule\n\\end{longtable}\n}")
        return tex_document(name, body, cites, landscape=len(cols) > 4).encode()

    # Markdown / HTML / PDF / Word / text share a Markdown table
    md_head = ["Paper", "Year"] + cols
    md_rows = []
    for i, (doc_id, r) in enumerate(rows, 1):
        d = docs[doc_id]
        paper = f"{d.title} ({short_cite(d)})" if with_citations else d.title
        md_rows.append([paper, str(d.year or "")] + [str(cells.get((doc_id, c), {}).get("value", "")) for c in cols])
    clean = lambda s: s.replace("|", "/").replace("\n", " ")
    md = "| " + " | ".join(md_head) + " |\n|" + "---|" * len(md_head) + "\n" + \
         "\n".join("| " + " | ".join(clean(c) for c in r) + " |" for r in md_rows)
    if with_citations and used:
        md += "\n\n## References\n\n" + "\n".join(f"- {format_reference(d, style)}"
                                                 for d in sorted(used, key=lambda d: short_cite(d).lower()))
    if fmt == "md":
        return (f"# {name}\n\n" + md).encode()
    if fmt == "txt":
        return text_from_md(md).encode()
    if fmt == "html":
        return html_document(name, html_from_md(md)).encode()
    if fmt == "pdf":
        return pdf_from_html(html_from_md(md), name, landscape=True)
    if fmt == "docx":
        from docx.enum.section import WD_ORIENT
        d = new_docx(name)
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = sec.page_height, sec.page_width
        docx_add_markdown(d, md, Citations({}, "none"))
        return docx_bytes(d)
    raise ValueError(fmt)


# ---- references export ------------------------------------------------------------------------------

def export_references(docs: list, fmt: str, style: str = "APA 7") -> bytes:
    if fmt == "bib":
        return bibtex(docs).encode()
    if fmt == "ris":
        return ris(docs).encode()
    if fmt == "csv":
        return references_csv(docs).encode("utf-8-sig")
    ordered = docs if style in NUMERIC_STYLES else sorted(docs, key=lambda d: short_cite(d).lower())
    entries = [format_reference(d, style, i if style in NUMERIC_STYLES else None) for i, d in enumerate(ordered, 1)]
    md = "\n\n".join(entries)
    title = f"References ({style})"
    if fmt == "md":
        return (f"# {title}\n\n" + md).encode()
    if fmt == "txt":
        return re.sub(r"\*", "", md).encode()
    if fmt == "docx":
        d = new_docx(title)
        for e in entries:
            _docx_inline(d.add_paragraph(), e, Citations({}, "none"))
        return docx_bytes(d)
    if fmt == "html":
        return html_document(title, html_from_md(md)).encode()
    if fmt == "pdf":
        return pdf_from_html(html_from_md(md), title)
    if fmt == "tex":
        cites = Citations({d.id: d for d in docs}, "numbered", style)
        cites.order = [d.id for d in docs]
        return tex_document(title, "\\nocite{*}\n", cites).encode()
    raise ValueError(fmt)


def zip_bytes(files: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data if isinstance(data, bytes) else data.encode())
    return buf.getvalue()
